# ------------------------- Import Libraries ------------------------- #

import argparse
import os
import pytesseract
import PIL.Image
from docx import Document
from fpdf import FPDF

# ------------------------- Function ------------------------- #

def image_to_str(input_file: str, output_file: str, lang: str, format: str):
    my_config_block_text = r"--psm 1 --oem 3"
    # My path is: H:\\Projects\\Python Files\\PyTesseract
    file_path = os.path.dirname(os.path.realpath(__file__))
    if lang:
        my_config_block_text = r"--psm 1 --oem 3 -l {}".format(lang)

    text = pytesseract.image_to_string(PIL.Image.open(input_file), config=my_config_block_text)
    
    if format == None:
        with open(file_path + "\\{}.{}".format(output_file, format), 'w', encoding='utf-8') as file:
            file.write(text)
    elif format == 'docx':
        document = Document()
        document.add_paragraph(text)
        document.save("{}.{}".format(output_file, format))
    elif format == 'pdf':
        pdf = FPDF()
        pdf.add_page()
        pdf.set_font('Arial', '',18)
        a = text.splitlines()
        pdf.write(10, text)
        # for i in range(len(a)):
        #     pdf.cell(40, 10, a)
        outp = output_file + "." + format
        pdf.output(outp)

    print("\n\nSuccess!! \n\nFile successfully generated on :\n\n>> {} <<\n\nFile name :\n [+] {}.{}".format(file_path, output_file, format))

# ------------------------- Main Section ------------------------- #

langhelp = """ 
Supported Languages :
    English: eng,
    Farsi: fas,
    Arabic: ara,
    French: fra,
    Italian: ita,
    Russian: rus,
"""
parser = argparse.ArgumentParser(description=langhelp)
parser.add_argument("-i", "--input", help="[+] Select Input File.")
parser.add_argument("-o", "--output", help="[+] Select Output File.")
parser.add_argument("-l", "--lang", help="[+] Select a Language.", default="eng")
parser.add_argument("-f", "--format", help="[+] Select a Format. [txt, docx, pdf]", default="txt")
args = parser.parse_args()


try:
    image_to_str(args.input, args.output, args.lang, args.format)

except:
    print("HEY .... Looks like you missed somethings.")



# ------------------------- Some Information About Tesseract Engine ------------------------- #
"""
Page Segmentation modes:
  0    Orientation and script detection (OSD) only.
  1    Automatic page segmentation with OSD.
  2    Automatic page segmentation, but no OSD, or OCR.
  3    Fully automatic page segmentation, but no OSD. (Default)
  4    Assume a single column of text of variable sizes.
  5    Assume a single uniform block of vertically aligned text.
  6    Assume a single uniform block of text.
  7    Trate the image as a single text line.
  8    Trate the image as a single word.
  9    Trate the image as a single word in a circle.
  10   Trate the image as a single character.
  11   Sparse text. Find as musch text as possible in no particular order.
  12   Sparse text with OSD.
  13   Raw line. Treat the image as a single text line,
			bypassing hack that are Tesseract-specific.
"""
"""
OCR Engine Mode
0    Legacy engine only.
1    Neural nets LSTM engine only.
2    Legacy + LSTM engines.
3    Default, based on what is available.
"""